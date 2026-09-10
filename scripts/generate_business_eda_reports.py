#!/usr/bin/env python3
"""Generate business EDA DOCX reports for the five non-FIRMS, non-population sources."""
from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

import geopandas as gpd
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import rasterio
from affine import Affine
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from rasterio.enums import Resampling
from rasterio.features import geometry_mask
from rasterio.mask import mask
from rasterio.windows import from_bounds

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = Path("/mnt/c/Users/USER/Downloads/FIRELINE_EDA_Business_Reports")
BLUE, GREEN, ORANGE, RED = "#155E75", "#0F766E", "#F4A261", "#E76F51"


def save_plot(path: Path) -> None:
    plt.tight_layout()
    plt.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close()


def shade(cell, color: str = "E7F3F5") -> None:
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def text(cell, value: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(value))
    run.bold = bold
    run.font.size = Pt(9)


def document(title: str, subtitle: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.6)
    section.left_margin, section.right_margin = Cm(1.7), Cm(1.7)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    for name, size in [("Heading 1", 17), ("Heading 2", 12)]:
        doc.styles[name].font.name = "Aptos Display"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(23, 78, 104)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(22), RGBColor(23, 78, 104)
    p = doc.add_paragraph(subtitle)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic, p.runs[0].font.size = True, Pt(10)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("FIRELINE — Business EDA | non-FIRMS dan non-populasi")
    r.font.size, r.font.color.rgb = Pt(8), RGBColor(95, 95, 95)
    return doc


def metrics(doc: Document, values: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ("Indikator", "Nilai")):
        text(cell, value, True)
        shade(cell, "C8E4E8")
    for key, value in values:
        row = table.add_row().cells
        text(row[0], key)
        text(row[1], value)


def feature_table(doc: Document, rows: list[tuple[str, str, str, str]]) -> None:
    """Add a complete source-schema table; it is intentionally not row-limited."""
    doc.add_heading("Arti kolom/data yang digunakan", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    headers = ("Kolom / data", "Bentuk data", "Arti sederhananya", "Dipakai untuk / catatan")
    for cell, value in zip(table.rows[0].cells, headers):
        text(cell, value, True)
        shade(cell, "C8E4E8")
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            text(cell, value)
    doc.add_paragraph()


def bullets(doc: Document, values: list[str]) -> None:
    for value in values:
        doc.add_paragraph(value, style="List Bullet")


def chart(doc: Document, path: Path, caption: str) -> None:
    doc.add_picture(str(path), width=Inches(6.55))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic, p.runs[0].font.size = True, Pt(8.5)


def insight(doc: Document, value: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell)
    p = cell.paragraphs[0]
    r = p.add_run("Business insight utama: ")
    r.bold, r.font.color.rgb = True, RGBColor(23, 78, 104)
    p.add_run(value)
    doc.add_paragraph()


def finish(doc: Document, output: Path, name: str) -> None:
    output.mkdir(parents=True, exist_ok=True)
    destination = output / name
    doc.save(destination)
    print(f"created {destination}")


def kalimantan() -> tuple[gpd.GeoDataFrame, str]:
    adm1 = gpd.read_file(ROOT / "data/raw/boundaries/idn_adm1.geojson")
    province_col = next(c for c in ("shapeName", "NAME_1", "name") if c in adm1.columns)
    result = adm1[adm1[province_col].astype(str).str.contains("Kalimantan", case=False, na=False)].copy()
    if result.empty:
        raise RuntimeError("Kalimantan was not found in ADM1 boundaries")
    return result, province_col


def clipped_raster(path: Path) -> np.ma.MaskedArray:
    provinces, _ = kalimantan()
    with rasterio.open(path) as source:
        shapes = [x.__geo_interface__ for x in provinces.geometry]
        image, _ = mask(source, shapes, crop=True, filled=False)
    return image[0]


def sampled_raster(
    path: Path,
    max_dimension: int = 1600,
    resampling: Resampling = Resampling.nearest,
) -> tuple[np.ma.MaskedArray, Affine]:
    """Return a bounded regular raster sample clipped to Kalimantan geometry."""
    provinces, _ = kalimantan()
    with rasterio.open(path) as source:
        window = from_bounds(*provinces.total_bounds, transform=source.transform)
        window = window.round_offsets().round_lengths()
        scale = min(1.0, max_dimension / max(window.width, window.height))
        height = max(1, round(window.height * scale))
        width = max(1, round(window.width * scale))
        data = source.read(
            1,
            window=window,
            out_shape=(height, width),
            resampling=resampling,
            masked=True,
        )
        transform = source.window_transform(window) * Affine.scale(
            window.width / width, window.height / height
        )
    inside = geometry_mask(
        [geometry.__geo_interface__ for geometry in provinces.geometry],
        out_shape=data.shape,
        transform=transform,
        invert=True,
    )
    return np.ma.array(data, mask=np.ma.getmaskarray(data) | ~inside), transform


def era5_report(output: Path, figures: Path) -> None:
    files = sorted((ROOT / "data/interim/era5_daily_parts").glob("era5_daily_*.parquet"))
    if not files:
        raise FileNotFoundError("ERA5 daily extracts are missing")
    days, grid_months = [], []
    for file in files:
        frame = pd.read_parquet(file, columns=["date", "grid_id", "rainfall_mm", "temperature_mean_c", "soil_moisture"])
        days.append(frame.groupby("date", as_index=False).mean(numeric_only=True))
        frame["month_start"] = pd.to_datetime(frame["date"]).dt.to_period("M").dt.to_timestamp()
        grid_months.append(frame.groupby(["grid_id", "month_start"], as_index=False).agg(
            rainfall_mm=("rainfall_mm", "sum"), soil=("soil_moisture", "mean")
        ))
    daily = pd.concat(days, ignore_index=True).drop_duplicates("date").sort_values("date")
    daily["date"] = pd.to_datetime(daily["date"])
    monthly = daily.set_index("date").resample("MS").agg(
        rainfall_mm=("rainfall_mm", "sum"), temperature_c=("temperature_mean_c", "mean"), soil=("soil_moisture", "mean")
    )
    climate = monthly.groupby(monthly.index.month).mean()
    names = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]
    driest, wettest = int(climate["rainfall_mm"].idxmin()), int(climate["rainfall_mm"].idxmax())
    soil_driest = int(climate["soil"].idxmin())
    climate["dryness_score"] = climate["rainfall_mm"].rank(pct=True) + climate["soil"].rank(pct=True)
    watch_months = climate.nsmallest(4, "dryness_score").index.tolist()
    grid_monthly = pd.concat(grid_months, ignore_index=True)
    grid_province = pd.read_parquet(
        ROOT / "data/processed/dim_grid.parquet", columns=["grid_id", "province"]
    ).drop_duplicates("grid_id")
    grid_monthly = grid_monthly.merge(grid_province, on="grid_id", how="inner", validate="many_to_one")
    if grid_monthly.empty or grid_monthly["province"].nunique() != 5:
        raise RuntimeError("ERA5 grid records could not be assigned to all five Kalimantan provinces")
    grid_monthly["rain_rank"] = grid_monthly.groupby("grid_id")["rainfall_mm"].rank(pct=True)
    grid_monthly["soil_rank"] = grid_monthly.groupby("grid_id")["soil"].rank(pct=True)
    grid_monthly["compound_dry"] = (grid_monthly["rain_rank"] <= .25) & (grid_monthly["soil_rank"] <= .25)
    compound_coverage = grid_monthly.groupby("month_start")["compound_dry"].mean().mul(100)
    widest_compound = pd.Timestamp(compound_coverage.idxmax())
    widest_share = float(compound_coverage.max())
    rain_soil_corr = float(monthly["rainfall_mm"].corr(monthly["soil"]))
    province_monthly = grid_monthly.groupby(["province", "month_start"], as_index=False).agg(
        rainfall_mm=("rainfall_mm", "mean"), soil=("soil", "mean")
    )
    province_climate = province_monthly.assign(month=lambda x: x.month_start.dt.month).groupby(
        ["province", "month"], as_index=False
    ).mean(numeric_only=True)
    province_climate["rain_rank"] = province_climate.groupby("province")["rainfall_mm"].rank(pct=True)
    province_climate["soil_rank"] = province_climate.groupby("province")["soil"].rank(pct=True)
    province_climate["dryness_score"] = province_climate["rain_rank"] + province_climate["soil_rank"]
    province_windows = province_climate.sort_values("dryness_score").groupby("province").head(3).groupby("province")["month"].apply(list)
    province_window_text = "; ".join(
        f"{province}: {', '.join(names[month - 1] for month in sorted(months))}"
        for province, months in province_windows.sort_index().items()
    )
    plt.figure(figsize=(9.6, 4.7))
    plt.bar(climate.index, climate["rainfall_mm"], color=BLUE, label="Curah hujan")
    plt.xticks(range(1, 13), names); plt.ylabel("Rata-rata hujan bulanan (mm)")
    plt.title("Klimatologi hujan untuk jendela kesiapsiagaan karhutla")
    for month in watch_months: plt.axvspan(month - .45, month + .45, color=RED, alpha=.10)
    plt.grid(axis="y", alpha=.25); plt.legend()
    save_plot(figures / "era5_rain.png")
    heatmap = compound_coverage.rename_axis("month_start").reset_index(name="share").assign(
        year=lambda x: x.month_start.dt.year, month=lambda x: x.month_start.dt.month
    ).pivot(index="year", columns="month", values="share").reindex(columns=range(1, 13))
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.1))
    axes[0].plot(climate.index, climate["soil"], marker="o", color=GREEN, label="Kelembapan tanah")
    axes[0].scatter([driest], [climate.loc[driest, "soil"]], color=RED, zorder=3, label="Bulan hujan minimum")
    axes[0].set(xticks=range(1, 13), xticklabels=names, ylabel="m³/m³", title="Siklus kelembapan tanah")
    axes[0].grid(alpha=.25); axes[0].legend(fontsize=8)
    image = axes[1].imshow(heatmap, cmap="YlOrRd", aspect="auto", vmin=0)
    axes[1].set(xticks=range(12), xticklabels=names, yticks=range(len(heatmap.index)), yticklabels=heatmap.index, title="Cakupan kekeringan majemuk (%)")
    axes[1].set_xlabel("Bulan"); axes[1].set_ylabel("Tahun")
    fig.colorbar(image, ax=axes[1], label="% grid pada kuartil kering lokal")
    save_plot(figures / "era5_compound_dryness.png")
    plt.figure(figsize=(7.5, 5.2))
    points = plt.scatter(monthly["rainfall_mm"], monthly["soil"], c=monthly.index.year, cmap="viridis", s=55, edgecolor="white", linewidth=.5)
    plt.colorbar(points, label="Tahun")
    plt.xlabel("Curah hujan bulanan regional (mm)"); plt.ylabel("Kelembapan tanah rata-rata (m³/m³)")
    plt.title("Hujan bulanan tidak selalu langsung memulihkan kelembapan tanah")
    plt.grid(alpha=.25)
    save_plot(figures / "era5_rain_soil_relation.png")
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2), sharex=True)
    for province, subset in province_climate.groupby("province", sort=True):
        subset = subset.sort_values("month")
        axes[0].plot(subset["month"], subset["rainfall_mm"], marker="o", linewidth=1.7, label=province)
        axes[1].plot(subset["month"], subset["soil"], marker="o", linewidth=1.7, label=province)
    axes[0].set(title="Musim hujan menurut provinsi", ylabel="Hujan rata-rata bulanan (mm)")
    axes[1].set(title="Musim kelembapan tanah menurut provinsi", ylabel="Kelembapan tanah (m³/m³)")
    for axis in axes:
        axis.set(xticks=range(1, 13), xticklabels=names, xlabel="Bulan")
        axis.grid(alpha=.25)
    axes[1].legend(fontsize=7.2, loc="best")
    save_plot(figures / "era5_province_climate.png")
    province_dryness = province_climate.pivot(index="province", columns="month", values="dryness_score").reindex(columns=range(1, 13))
    plt.figure(figsize=(10.4, 3.8))
    image = plt.imshow(province_dryness, cmap="YlOrRd", aspect="auto")
    plt.xticks(range(12), names); plt.yticks(range(len(province_dryness.index)), province_dryness.index)
    plt.xlabel("Bulan"); plt.ylabel("Provinsi"); plt.title("Kalender kekeringan relatif per provinsi")
    plt.colorbar(image, label="Skor gabungan peringkat hujan + kelembapan; lebih rendah = lebih kering")
    save_plot(figures / "era5_province_dryness.png")
    doc = document("ERA5-Land: iklim dan kesiapan operasi", "EDA bisnis Kalimantan | tanpa NASA FIRMS maupun data populasi")
    doc.add_heading("Tentang dataset", 1)
    doc.add_paragraph("ERA5-Land adalah reanalisis permukaan dari Copernicus Climate Change Service. Data mentah proyek mencakup 43 berkas bulanan (Januari 2023–Juli 2026) dengan curah hujan total, suhu 2 m, dan kelembapan tanah lapisan pertama. Analisis merangkum rata-rata regional sekaligus membandingkan kondisi bulanan pada setiap grid. ERA5 tidak mendeteksi api, tetapi dapat membaca prasyarat pengeringan bahan bakar vegetasi dan gambut sebelum kebakaran berkembang.")
    feature_table(doc, [
        ("valid_time", "Tanggal dan jam", "Kapan kondisi cuaca direkam oleh ERA5", "Dikelompokkan menjadi hari, bulan, dan musim; waktu memakai UTC"),
        ("latitude", "Koordinat", "Posisi utara–selatan sebuah sel peta cuaca", "Menentukan lokasi grid; bukan nama kabupaten"),
        ("longitude", "Koordinat", "Posisi timur–barat sebuah sel peta cuaca", "Digabung dengan latitude untuk menempatkan grid di Kalimantan"),
        ("tp", "Curah hujan", "Total hujan pada sel cuaca", "Diubah ke milimeter untuk membaca periode hujan rendah"),
        ("t2m", "Suhu udara", "Suhu udara sekitar dua meter dari permukaan", "Memberi konteks pengeringan; bukan suhu titik lapangan"),
        ("swvl1", "Kelembapan tanah", "Banyaknya air di lapisan tanah paling atas", "Dipakai bersama hujan untuk melihat kekeringan tanah relatif"),
        ("grid_id (buatan proyek)", "Kode lokasi", "Nama singkat untuk tiap sel grid analisis", "Membandingkan satu lokasi dengan riwayatnya sendiri, lalu mengelompokkan ke provinsi"),
    ])
    metrics(doc, [("Cakupan", f"{daily.date.min():%d %b %Y} – {daily.date.max():%d %b %Y}"), ("Ringkasan hari", f"{len(daily):,}"), ("Berkas mentah", f"{len(files)} bulan"), ("Jendela siaga regional", ", ".join(names[x-1] for x in watch_months)), ("Jendela lokal (3 bulan)", province_window_text), ("Kekeringan majemuk terluas", f"{widest_compound:%b %Y} ({widest_share:.1f}% grid)"), ("Korelasi hujan–kelembapan bulanan", f"r = {rain_soil_corr:.2f}")])
    doc.add_heading("Exploratory Data Analysis", 1)
    chart(doc, figures / "era5_rain.png", "Gambar 1. Arsiran merah menunjukkan empat bulan dengan kombinasi hujan dan kelembapan tanah paling rendah secara klimatologis.")
    doc.add_paragraph(f"Bulan {', '.join(names[x-1] for x in watch_months)} membentuk jendela siaga relatif. {names[driest-1]} memiliki rata-rata hujan terendah, sementara minimum kelembapan tanah terjadi pada {names[soil_driest-1]}. Perbedaan waktunya penting: kesiapsiagaan tidak boleh dimulai hanya setelah hujan telah sangat rendah, karena penyimpanan air tanah dapat terus menurun atau pulih dengan ritme yang berbeda.")
    chart(doc, figures / "era5_compound_dryness.png", "Gambar 2. Kekeringan majemuk: grid yang berada pada kuartil hujan dan kuartil kelembapan tanah terendah dibanding riwayatnya sendiri.")
    chart(doc, figures / "era5_province_climate.png", "Gambar 3. Rata-rata musim hujan dan kelembapan tanah pada grid yang berada di masing-masing provinsi Kalimantan.")
    doc.add_paragraph(f"Pola antarprovinsi tidak identik. Tiga bulan dengan kombinasi hujan dan kelembapan tanah paling rendah relatif terhadap riwayat provinsinya adalah: {province_window_text}. Karena itu kalender regional berguna untuk koordinasi bersama, tetapi penempatan pompa, patroli, dan audit kanal harus mengikuti jendela lokal—bukan satu tanggal yang dipaksakan untuk seluruh pulau.")
    chart(doc, figures / "era5_province_dryness.png", "Gambar 4. Kalender relatif per provinsi; nilai rendah menunjukkan kombinasi hujan dan kelembapan tanah yang lebih kering dibanding bulan lain di provinsi yang sama.")
    doc.add_paragraph("Peta panas ini memisahkan masalah kalender dari masalah skala wilayah. Dua provinsi dapat sama-sama berada pada musim kering regional, tetapi intensitas relatif dan waktunya tidak identik. Command centre dapat memakai matriks ini untuk menjadwalkan eskalasi logistik berurutan, sehingga armada pompa dan tim audit tidak harus ditempatkan pada puncak kesiapan yang sama di seluruh provinsi.")
    doc.add_heading("Implikasi terhadap karhutla Kalimantan", 1)
    doc.add_paragraph(f"Kekeringan majemuk paling luas pada {widest_compound:%B %Y}, ketika {widest_share:.1f}% grid berada sekaligus pada kuartil hujan dan kelembapan tanah terendahnya. Ini bukan jumlah area terbakar dan bukan probabilitas kebakaran. Namun, ketika cakupan kondisi kering meluas, gangguan kecil seperti pembakaran lahan, kanal yang kehilangan muka air, atau api permukaan lebih berpotensi membutuhkan respons lintas wilayah daripada respons kabupaten yang terisolasi.")
    doc.add_paragraph(f"Korelasi bulanan hujan dengan kelembapan tanah adalah r={rain_soil_corr:.2f}. Hubungan yang tidak sempurna ini memperlihatkan mengapa satu indikator hujan saja tidak cukup untuk manajemen karhutla: kondisi air tersimpan di tanah—terutama pada lanskap gambut—perlu dipantau terpisah untuk menentukan kapan pencegahan air dan patroli diperketat.")
    chart(doc, figures / "era5_rain_soil_relation.png", "Gambar 5. Hubungan antarbulan antara hujan dan kelembapan tanah; warna menunjukkan tahun observasi.")
    doc.add_paragraph("Sebaran titik pada grafik memperlihatkan bahwa bulan dengan hujan yang mirip dapat memiliki kelembapan tanah yang berbeda. Artinya, keputusan menurunkan atau menaikkan status kesiapsiagaan tidak cukup hanya memakai akumulasi hujan; riwayat pengeringan tanah perlu ikut dibaca untuk menghindari penurunan kesiapan yang terlalu dini.")
    insight(doc, f"Gunakan {', '.join(names[x-1] for x in watch_months)} sebagai fase *pre-positioning*: audit kanal dan pompa, penempatan bahan bakar/logistik, pemutakhiran daftar kontak desa, serta pengawasan pembukaan lahan. Saat cakupan kekeringan majemuk naik lintas grid seperti pada {widest_compound:%b %Y}, aktifkan koordinasi lintas provinsi lebih awal—bukan menunggu laporan api membesar.")
    doc.add_heading("Rekomendasi keputusan", 1)
    bullets(doc, ["Bangun *two-stage trigger*: kalender pre-positioning dari empat bulan siaga, lalu trigger harian/mingguan dari hujan dan kelembapan tanah lokal.", "Jadwalkan audit hidrologi gambut dan kesiapan sumber air sebelum kelembapan tanah mencapai fase minimum, bukan setelah muncul krisis air.", "Gunakan peta cakupan kekeringan majemuk untuk menentukan kapan command centre perlu menyatukan logistik lintas provinsi.", "Turunkan analisis ke provinsi/kabupaten dan overlay-kan dengan gambut, tutupan lahan, akses, serta aset sebelum membuat ambang dan penugasan lokal."])
    doc.add_heading("Batasan interpretasi", 1)
    doc.add_paragraph("ERA5-Land adalah produk reanalisis berbasis grid; nilainya tidak menggantikan alat ukur lapangan. Laporan ini tidak menilai fitur maupun kinerja model machine learning.")
    finish(doc, output, "01_ERA5_Land_Climate_Operations_EDA.docx")


def landcover_report(output: Path, figures: Path) -> None:
    raster, sample_transform = sampled_raster(ROOT / "data/raw/mapbiomas/landcover_2022_collection2.tif")
    values = raster.compressed().astype("int64"); values = values[values != 0]
    codes, counts = np.unique(values, return_counts=True)
    summary = pd.DataFrame({"code": codes, "pixels": counts})
    summary["share_pct"] = summary.pixels / summary.pixels.sum() * 100
    summary = summary.sort_values("share_pct", ascending=False)
    top, concentration = summary.head(8), summary.head(3).share_pct.sum()
    plt.figure(figsize=(9.5, 4.7))
    plt.bar(top.code.astype(str), top.share_pct, color=[BLUE, GREEN, "#E9C46A", ORANGE, RED, "#6D597A", "#457B9D", "#8AB17D"])
    plt.xlabel("Kode kelas MapBiomas pada raster"); plt.ylabel("Pangsa piksel valid (%)")
    plt.title("Delapan kelas tutupan lahan terbesar di Kalimantan (2022)"); plt.grid(axis="y", alpha=.25)
    save_plot(figures / "landcover_share.png")
    step = max(1, int(max(raster.shape) / 900))
    plt.figure(figsize=(7, 5.3)); image = plt.imshow(raster[::step, ::step], cmap="tab20", interpolation="nearest")
    plt.colorbar(image, label="Kode kelas MapBiomas"); plt.title("Mosaik tutupan lahan Kalimantan (2022)"); plt.axis("off")
    save_plot(figures / "landcover_map.png")
    provinces, province_col = kalimantan()
    valid_cells = ~np.ma.getmaskarray(raster)
    province_rows = []
    for _, province in provinces.iterrows():
        inside = geometry_mask([province.geometry.__geo_interface__], out_shape=raster.shape, transform=sample_transform, invert=True)
        province_values = np.asarray(raster.data[inside & valid_cells]).astype("int64")
        province_values = province_values[province_values != 0]
        province_codes, province_counts = np.unique(province_values, return_counts=True)
        dominant_index = int(np.argmax(province_counts))
        province_rows.append((str(province[province_col]), int(province_codes[dominant_index]), province_counts[dominant_index] / province_counts.sum() * 100, len(province_codes)))
    province_mix = pd.DataFrame(province_rows, columns=["province", "dominant_code", "dominant_share_pct", "class_count"]).sort_values("dominant_share_pct")
    plt.figure(figsize=(8.6, 4.8)); bars = plt.barh(province_mix.province, province_mix.dominant_share_pct, color=GREEN)
    for bar, code in zip(bars, province_mix.dominant_code): plt.text(bar.get_width() + .3, bar.get_y() + bar.get_height()/2, f"kode {code}", va="center", fontsize=8)
    plt.xlabel("Pangsa kode tutupan lahan dominan (%)"); plt.title("Konsentrasi lanskap dominan per provinsi (sampel 2022)"); plt.grid(axis="x", alpha=.25)
    save_plot(figures / "landcover_province_concentration.png")
    temporal = pd.read_parquet(ROOT / "data/interim/mapbiomas_province_year_features.parquet")
    temporal = temporal.loc[temporal.source_year.isin([2023, 2024])].copy()
    required = {"province", "source_year", "lc_natural_share_pct", "lc_anthropic_share_pct", "lc_peat_swamp_share_pct"}
    if len(temporal) != 10 or not required.issubset(temporal.columns):
        raise RuntimeError("MapBiomas 2023–2024 province land-cover statistics are incomplete")
    trend = temporal.pivot(index="province", columns="source_year", values=["lc_natural_share_pct", "lc_anthropic_share_pct", "lc_peat_swamp_share_pct"])
    change = pd.DataFrame({
        "natural_change_pp": trend[("lc_natural_share_pct", 2024)] - trend[("lc_natural_share_pct", 2023)],
        "anthropic_change_pp": trend[("lc_anthropic_share_pct", 2024)] - trend[("lc_anthropic_share_pct", 2023)],
        "peat_swamp_change_pp": trend[("lc_peat_swamp_share_pct", 2024)] - trend[("lc_peat_swamp_share_pct", 2023)],
    }).sort_values("natural_change_pp")
    largest_natural_decline = change.iloc[0]
    plt.figure(figsize=(9.2, 4.9))
    y = np.arange(len(change)); height = .26
    plt.barh(y - height, change["natural_change_pp"], height=height, color=GREEN, label="Perubahan natural")
    plt.barh(y, change["anthropic_change_pp"], height=height, color=ORANGE, label="Perubahan antropik")
    plt.barh(y + height, change["peat_swamp_change_pp"], height=height, color="#6D597A", label="Perubahan rawa gambut")
    plt.yticks(y, change.index); plt.axvline(0, color="#555555", linewidth=.8)
    plt.xlabel("Perubahan pangsa 2024 vs 2023 (poin persentase)"); plt.title("Perubahan komposisi tutupan lahan per provinsi")
    plt.grid(axis="x", alpha=.25); plt.legend(fontsize=8)
    save_plot(figures / "landcover_province_change_2023_2024.png")
    latest_mix = temporal.loc[temporal.source_year.eq(2024), ["province", "lc_natural_share_pct", "lc_anthropic_share_pct", "lc_peat_swamp_share_pct"]].set_index("province").sort_values("lc_anthropic_share_pct")
    fig, ax = plt.subplots(figsize=(9.2, 4.8))
    ax.barh(latest_mix.index, latest_mix["lc_natural_share_pct"], color=GREEN, label="Natural")
    ax.barh(latest_mix.index, latest_mix["lc_anthropic_share_pct"], left=latest_mix["lc_natural_share_pct"], color=ORANGE, label="Antropik")
    ax.set(xlabel="Pangsa luas provinsi (%)", title="Portofolio lanskap per provinsi (2024)", xlim=(0, 100))
    ax.grid(axis="x", alpha=.25); ax.legend()
    save_plot(figures / "landcover_province_portfolio_2024.png")
    most_anthropic = latest_mix.iloc[-1]
    doc = document("MapBiomas: komposisi tutupan lahan", "EDA bisnis Kalimantan | baseline raster 2022 dan statistik resmi 2023–2024")
    doc.add_heading("Tentang dataset", 1)
    doc.add_paragraph("GeoTIFF MapBiomas Indonesia Collection 2 tahun 2022 menyimpan klasifikasi tutupan/penggunaan lahan berbasis piksel. Analisis dibatasi pada Kalimantan memakai geoBoundaries ADM1. Agar aman dijalankan pada laptop, statistik raster 2022 dihasilkan dari sampel reguler berukuran terbatas; nilai nol diperlakukan sebagai non-data. Untuk pembaruan, tabel statistik resmi MapBiomas Collection 4.1 memberi komposisi provinsi 2023 dan 2024. Tidak ada produk tutupan lahan tahunan 2025–2026 yang dapat diunduh secara resmi pada saat pengambilan, sehingga laporan tidak mengarang tren dua tahun tersebut.")
    feature_table(doc, [
        ("Band 1 — class code", "Kode angka per piksel", "Kode jenis tutupan/penggunaan lahan pada satu kotak kecil peta", "Menunjukkan komposisi dan mosaik lanskap; nama kelas harus dikonfirmasi lewat legenda resmi MapBiomas"),
        ("Nilai 0", "Kode khusus", "Piksel tanpa data yang layak dianalisis", "Dikeluarkan dari perhitungan proporsi"),
        ("Posisi piksel", "Lokasi pada peta", "Di mana setiap kode kelas berada", "Memotong data ke Kalimantan dan merangkum per provinsi"),
        ("Transform / resolusi raster", "Info ukuran peta", "Ukuran, arah, dan susunan kotak-kotak peta", "Memastikan sampling aman untuk laptop; bukan indikator risiko"),
        ("CRS raster", "Sistem koordinat", "Cara GeoTIFF menyimpan posisi bumi", "Menyelaraskan raster dengan batas provinsi"),
        ("Tahun peta", "Tahun", "Raster baseline Collection 2 dibuat untuk 2022", "Memberi konteks lanskap historis, bukan kondisi pasti tahun 2026"),
        ("Statistik Coverage C4.1", "Tabel luas per kelas", "Statistik resmi komposisi lahan per provinsi untuk 2023 dan 2024", "Membandingkan pangsa natural, antropik, dan rawa gambut tanpa mengarang data 2025–2026"),
    ])
    metrics(doc, [("Baseline raster", "2022"), ("Statistik provinsi resmi", "2023–2024 (Collection 4.1)"), ("Sel sampel valid", f"{len(values):,}"), ("Kode kelas teramati", f"{len(summary):,}"), ("Konsentrasi tiga kode terbesar", f"{concentration:.1f}%"), ("Kode paling dominan", f"{int(summary.iloc[0].code)} ({summary.iloc[0].share_pct:.1f}%)")])
    doc.add_heading("Exploratory Data Analysis", 1)
    chart(doc, figures / "landcover_share.png", "Gambar 1. Pangsa dihitung dari sel sampel valid, bukan dari luas administratif resmi.")
    doc.add_paragraph(f"Tiga kode terbesar mencakup {concentration:.1f}% sel sampel valid. Komposisi yang terkonsentrasi berarti strategi pencegahan dapat dipusatkan pada sejumlah kecil tipe lanskap dominan, dengan kelas kecil tetap diperiksa bila berada di lokasi bernilai tinggi.")
    chart(doc, figures / "landcover_map.png", "Gambar 2. Mosaik untuk melihat heterogenitas lanskap; warna adalah kode raster, bukan tingkat risiko.")
    chart(doc, figures / "landcover_province_concentration.png", "Gambar 3. Pangsa kode dominan per provinsi; label menunjukkan kode MapBiomas, bukan nama kelas yang ditebak.")
    doc.add_paragraph("EDA tambahan ini membedakan provinsi dengan lanskap yang relatif terkonsentrasi dari provinsi yang lebih beragam. Pada provinsi yang sangat terkonsentrasi, paket pencegahan dapat dirancang lebih fokus pada satu tipe lanskap utama. Pada provinsi yang pangsa dominannya lebih rendah, desain program perlu lebih modular karena strategi bahan bakar, akses, dan engagement pengelola kemungkinan tidak homogen.")
    chart(doc, figures / "landcover_province_change_2023_2024.png", "Gambar 4. Perubahan pangsa kelas dari statistik provinsi resmi MapBiomas Collection 4.1; ini bukan perubahan yang dihitung dari sampel raster 2022.")
    doc.add_paragraph(f"{largest_natural_decline.name} menunjukkan penurunan pangsa natural terbesar pada 2024 dibanding 2023 ({largest_natural_decline.natural_change_pp:.2f} poin persentase), dengan perubahan pangsa antropik {largest_natural_decline.anthropic_change_pp:+.2f} poin. Skala perubahan satu tahun ini kecil dan tidak menyatakan sebab atau legalitas konversi, tetapi cukup untuk menjadikan provinsi tersebut prioritas verifikasi lokasi perubahan, pengelola, dan kesiapan pencegahan pada lanskap transisi.")
    chart(doc, figures / "landcover_province_portfolio_2024.png", "Gambar 5. Komposisi natural–antropik pada setiap provinsi dari statistik MapBiomas 2024; segmen tidak menyatakan tingkat bahaya api.")
    doc.add_paragraph(f"{most_anthropic.name} memiliki pangsa lanskap antropik terbesar pada 2024 ({most_anthropic.lc_anthropic_share_pct:.1f}%). Ini membantu membedakan fokus pelibatan: provinsi dengan pangsa antropik lebih besar memerlukan pemetaan pengelola dan SOP pembakaran yang lebih granular, sedangkan lanskap dominan natural memerlukan strategi akses, perlindungan habitat, dan pencegahan koridor yang berbeda.")
    doc.add_heading("Implikasi terhadap karhutla Kalimantan", 1)
    doc.add_paragraph("Tutupan lahan tidak menyatakan bahwa api telah terjadi, tetapi menentukan konteks bahan bakar, kontinuitas lanskap, akses, dan pihak yang perlu diajak mencegah pembakaran. Strategi pada bentang hutan/vegetasi alami, mosaik budidaya, kawasan basah, dan area terbangun tidak boleh disamakan: jalur pencegahan, pesan komunikasi, serta aset responsnya berbeda.")
    doc.add_paragraph("Konsentrasi kelas dominan membuat program berbasis lanskap lebih efisien daripada kampanye seragam. Sebaliknya, mosaik kelas pada peta menandakan area transisi yang perlu diprioritaskan untuk koordinasi multi-pengelola, karena satu kejadian dapat melintasi tipe tutupan dan kewenangan yang berbeda.")
    insight(doc, "Gunakan peta ini untuk membagi portofolio pencegahan karhutla: strategi bahan bakar dan akses untuk lanskap dominan; strategi kemitraan serta kepatuhan pembakaran untuk mosaik penggunaan lahan; dan perlindungan muka air untuk kelas yang beririsan dengan gambut. Pemaknaan kelas wajib memakai legenda MapBiomas Collection 2 resmi sebelum SOP per kelas ditetapkan.")
    doc.add_heading("Rekomendasi keputusan", 1)
    bullets(doc, ["Simpan legenda resmi Collection 2 ke pipeline sebelum memberi label kelas pada dashboard atau SOP.", "Gunakan komposisi tutupan lahan untuk membagi portofolio program per lanskap, bukan sebagai skor bahaya kebakaran.", "Bandingkan peta lintas tahun untuk mengidentifikasi lanskap yang berubah cepat dan butuh engagement lebih awal."])
    doc.add_heading("Batasan interpretasi", 1)
    doc.add_paragraph("Kelas piksel adalah klasifikasi penginderaan jauh, bukan verifikasi kepemilikan maupun kondisi lapangan. Luas 2022 di sini berupa proporsi piksel sampel dan bukan klaim hektare resmi; seri 2023–2024 menggunakan statistik resmi per provinsi. Ketiadaan produk resmi 2025–2026 berarti nilai terakhir tidak boleh diekstrapolasikan sebagai observasi baru.")
    finish(doc, output, "02_MapBiomas_Land_Cover_EDA.docx")


def burned_report(output: Path, figures: Path) -> None:
    provinces, province_col = kalimantan(); path = ROOT / "data/raw/mapbiomas/burned_area_2024_collection2.tif"
    raster, sample_transform = sampled_raster(path, max_dimension=2000, resampling=Resampling.average)
    rows = []
    valid = ~np.ma.getmaskarray(raster)
    for _, province in provinces.iterrows():
        inside = geometry_mask([province.geometry.__geo_interface__], out_shape=raster.shape, transform=sample_transform, invert=True)
        values = np.asarray(raster.data[inside & valid]); values = values[np.isfinite(values)]
        # The annual layer is binary; average resampling retains burned-pixel share per block.
        rows.append((str(province[province_col]), len(values), float(np.clip(values, 0, 1).mean() * 100) if len(values) else 0))
    summary = pd.DataFrame(rows, columns=["province", "valid_pixels", "burned_pixel_share_pct"]).sort_values("burned_pixel_share_pct", ascending=False)
    top = summary.iloc[0]
    plt.figure(figsize=(9.4, 4.7)); plt.barh(summary.province, summary.burned_pixel_share_pct, color=RED); plt.gca().invert_yaxis()
    plt.xlabel("Pangsa piksel terbakar dari piksel valid (%)"); plt.title("Jejak area terbakar MapBiomas per provinsi (2024)"); plt.grid(axis="x", alpha=.25)
    save_plot(figures / "burned_bars.png")
    plt.figure(figsize=(7, 5.3)); image = plt.imshow(raster, cmap="Reds", interpolation="nearest", vmin=0, vmax=1); plt.colorbar(image, label="Pangsa terbakar pada blok raster"); plt.title("Jejak area terbakar MapBiomas (2024)"); plt.axis("off")
    save_plot(figures / "burned_map.png")
    intensity = np.clip(raster.compressed().astype(float), 0, 1)
    intensity = np.sort(intensity[intensity > 0])[::-1]
    cumulative = np.cumsum(intensity) / intensity.sum()
    block_share = np.arange(1, len(intensity) + 1) / len(intensity)
    top_tenth = float(cumulative[max(0, int(np.ceil(len(intensity) * .10)) - 1)] * 100)
    plt.figure(figsize=(7.5, 5.0)); plt.plot(block_share * 100, cumulative * 100, color=RED, linewidth=2, label="Jejak terbakar kumulatif")
    plt.plot([0, 100], [0, 100], color="#777777", linestyle="--", label="Sebaran merata")
    plt.xlabel("Blok dengan observasi terbakar, diurutkan dari terbesar (%)"); plt.ylabel("Jejak terbakar kumulatif (%)")
    plt.title("Konsentrasi spasial footprint terbakar (2024)"); plt.grid(alpha=.25); plt.legend()
    save_plot(figures / "burned_concentration.png")
    temporal = pd.read_parquet(ROOT / "data/interim/mapbiomas_province_year_features.parquet")
    annual = temporal.loc[temporal.source_year.isin([2023, 2024]), ["province", "source_year", "burned_share_pct"]].copy()
    if len(annual) != 10 or annual.duplicated(["province", "source_year"]).any():
        raise RuntimeError("MapBiomas annual burned-area statistics are incomplete")
    annual_pivot = annual.pivot(index="province", columns="source_year", values="burned_share_pct").sort_values(2023, ascending=False)
    annual_change = (annual_pivot[2024] - annual_pivot[2023]).sort_values()
    fig, ax = plt.subplots(figsize=(9.2, 4.9))
    y = np.arange(len(annual_pivot)); height = .36
    ax.barh(y - height / 2, annual_pivot[2023], height=height, color=ORANGE, label="2023")
    ax.barh(y + height / 2, annual_pivot[2024], height=height, color=RED, label="2024")
    ax.set(yticks=y, yticklabels=annual_pivot.index, xlabel="Pangsa area terbakar dari piksel valid (%)", title="Perbandingan footprint terbakar tahunan per provinsi")
    ax.invert_yaxis(); ax.grid(axis="x", alpha=.25); ax.legend()
    save_plot(figures / "burned_annual_2023_2024.png")
    monthly_rows = []
    for month in range(1, 7):
        monthly_path = ROOT / f"data/raw/mapbiomas/burned_area_monthly_2026_{month:02d}.tif"
        month_raster, month_transform = sampled_raster(monthly_path, max_dimension=1200, resampling=Resampling.average)
        month_valid = ~np.ma.getmaskarray(month_raster)
        for _, province in provinces.iterrows():
            inside = geometry_mask([province.geometry.__geo_interface__], out_shape=month_raster.shape, transform=month_transform, invert=True)
            month_values = np.asarray(month_raster.data[inside & month_valid])
            month_values = month_values[np.isfinite(month_values)]
            monthly_rows.append((str(province[province_col]), month, float(np.clip(month_values, 0, 1).mean() * 100) if len(month_values) else 0.0))
    monthly_2026 = pd.DataFrame(monthly_rows, columns=["province", "month", "burned_pixel_share_pct"])
    peak_month = monthly_2026.loc[monthly_2026.burned_pixel_share_pct.idxmax()]
    fig, ax = plt.subplots(figsize=(9.2, 4.8))
    for province, subset in monthly_2026.groupby("province", sort=True):
        ax.plot(subset.month, subset.burned_pixel_share_pct, marker="o", linewidth=1.8, label=province)
    ax.set(xticks=range(1, 7), xticklabels=["Jan", "Feb", "Mar", "Apr", "Mei", "Jun"], xlabel="Bulan 2026", ylabel="Pangsa area terbakar bulanan (%)", title="Pemantauan MapBiomas Fire per provinsi (Jan–Jun 2026)")
    ax.grid(alpha=.25); ax.legend(fontsize=7.2)
    save_plot(figures / "burned_monthly_2026.png")
    doc = document("MapBiomas: jejak area terbakar", "EDA bisnis Kalimantan | annual 2023–2024 dan pemantauan Jan–Jun 2026")
    doc.add_heading("Tentang dataset", 1)
    doc.add_paragraph("GeoTIFF MapBiomas Indonesia Fire Collection 2 memetakan area yang terdeteksi terbakar. EDA merangkum produk annual 2023 dan 2024, serta produk pemantauan bulanan Januari–Juni 2026, pada setiap provinsi Kalimantan. Agregasi rata-rata menjaga perhitungan aman dijalankan pada laptop; metriknya adalah proporsi area, bukan estimasi hektare resmi. Produk annual 2025 dan 2026 maupun produk bulanan 2025 belum tersedia resmi saat pengambilan, sehingga tidak ada nilai yang diisi atau diproyeksikan untuk celah tersebut.")
    feature_table(doc, [
        ("Band 1 — burned-area value", "Nilai per piksel", "Tanda apakah permukaan pada kotak peta teramati terbakar", "Dipakai untuk melihat jejak/footprint area terbakar, bukan panas atau intensitas api"),
        ("Nilai 0", "Nilai raster", "Tidak ada area terbakar yang teramati pada piksel", "Tetap dihitung sebagai pembanding cakupan wilayah"),
        ("Nilai positif", "Nilai raster", "Ada area terbakar yang teramati pada piksel", "Menandai jejak kebakaran; tidak memberi tahu penyebab atau waktu penyalaan"),
        ("Posisi piksel", "Lokasi pada peta", "Di mana footprint terbakar berada", "Membandingkan provinsi dan menemukan klaster untuk audit lapangan"),
        ("Transform / resolusi raster", "Info ukuran peta", "Ukuran dan susunan kotak-kotak peta", "Dipakai saat merangkum secara aman; bukan ukuran luas resmi"),
        ("Tahun produk annual", "Tahun", "Peta jejak terbakar untuk satu tahun penuh: 2023 dan 2024", "Untuk evaluasi pascamusim dan perencanaan pencegahan berikutnya"),
        ("Bulan produk monitoring", "Bulan", "Peta pemantauan bulanan yang tersedia: Januari–Juni 2026", "Untuk membaca situasi terkini; belum boleh dianggap angka tahunan atau dijumlahkan mentah"),
    ])
    metrics(doc, [("Produk annual", "2023 dan 2024"), ("Monitoring terkini", "Jan–Jun 2026; belum satu tahun penuh"), ("Provinsi", str(len(summary))), ("Blok raster valid (2024)", f"{summary.valid_pixels.sum():,}"), ("Pangsa terbakar tertinggi (2024)", f"{top.province} ({top.burned_pixel_share_pct:.2f}%)"), ("Puncak bulanan 2026", f"{peak_month.province}, {['Jan','Feb','Mar','Apr','Mei','Jun'][int(peak_month.month)-1]} ({peak_month.burned_pixel_share_pct:.3f}%)")])
    doc.add_heading("Exploratory Data Analysis", 1)
    chart(doc, figures / "burned_bars.png", "Gambar 1. Perbandingan berdasarkan pangsa area terbakar pada blok raster valid tiap provinsi.")
    doc.add_paragraph(f"{top.province} memiliki pangsa piksel terbakar tertinggi pada data 2024 ({top.burned_pixel_share_pct:.2f}%). Perbedaan antardaerah adalah dasar untuk evaluasi pascakejadian dan pembelajaran kesiapan, bukan penilaian kinerja institusi tanpa konteks kejadian dan tutupan lahan.")
    chart(doc, figures / "burned_map.png", "Gambar 2. Merah menunjukkan pangsa area terbakar dalam blok raster, bukan intensitas atau waktu kejadian.")
    chart(doc, figures / "burned_concentration.png", "Gambar 3. Konsentrasi footprint pada blok raster yang sudah memiliki observasi terbakar.")
    doc.add_paragraph(f"Sepuluh persen blok dengan footprint terbakar terbesar menyumbang {top_tenth:.1f}% dari total footprint pada blok yang terdeteksi terbakar. Jika konsentrasi tinggi, audit pascakejadian, pemulihan akses, dan verifikasi pengelola dapat dimulai dari klaster terbesar tanpa menyamakan seluruh wilayah sebagai prioritas yang setara.")
    chart(doc, figures / "burned_annual_2023_2024.png", "Gambar 4. Perbandingan produk annual MapBiomas Fire Collection 2; perbedaan bukan estimasi perubahan hektare resmi.")
    largest_reduction = annual_change.index[0]
    doc.add_paragraph(f"Perbandingan tahunan memperlihatkan perubahan yang sangat berbeda antarprovinsi. Penurunan pangsa terbesar dari 2023 ke 2024 berada di {largest_reduction} ({annual_change.iloc[0]:.2f} poin persentase). Ini adalah sinyal untuk mengevaluasi apakah perubahan berasal dari kondisi cuaca, intervensi, pelaporan/observabilitas, atau karakter lanskap—bukan kesimpulan otomatis bahwa risiko telah hilang.")
    chart(doc, figures / "burned_monthly_2026.png", "Gambar 5. Produk monitoring bulanan yang tersedia hanya sampai Juni 2026; tiap titik adalah pangsa footprint bulan tersebut dan tidak boleh dijumlahkan menjadi angka annual tanpa pemeriksaan overlap piksel.")
    doc.add_paragraph(f"Pada monitoring Januari–Juni 2026, puncak pangsa bulanan pada sampel berada di {peak_month.province} pada {['Januari','Februari','Maret','April','Mei','Juni'][int(peak_month.month)-1]} ({peak_month.burned_pixel_share_pct:.3f}%). Gunakan sebagai daftar situasi terkini untuk verifikasi lapangan dan kesiapan respons, sambil tetap menjaga label *year-to-date* karena Juli–Desember belum termasuk.")
    doc.add_heading("Implikasi terhadap karhutla Kalimantan", 1)
    doc.add_paragraph("Area terbakar tahunan adalah indikator dampak yang bersifat tertinggal: ia tidak dapat memberi peringatan sebelum api muncul, tetapi sangat berguna untuk memutuskan lokasi pemulihan, inspeksi akses, dan evaluasi kegagalan pencegahan. Peta jejak memungkinkan unit operasi menutup siklus dari respons menjadi pembelajaran, bukan mengulang penyebaran sumber daya berdasarkan persepsi semata.")
    doc.add_paragraph("Untuk musim berikutnya, wilayah footprint perlu dicek bersama kondisi gambut, tutupan lahan, jalur air, dan pemegang hak/pengelola. Kombinasi ini membedakan lokasi yang perlu pemulihan hidrologis, pengurangan bahan bakar, pengawasan pembakaran, atau perbaikan akses respons.")
    insight(doc, f"Mulai review pascamusim dan anggaran pemulihan dari {top.province}, lalu audit setiap klaster footprint terhadap gambut, akses pemadaman, dan pengelola lahannya. Perlakukan lokasi tersebut sebagai daftar tindakan pencegahan tahun berikutnya—bukan hanya arsip kejadian 2024.")
    doc.add_heading("Rekomendasi keputusan", 1)
    bullets(doc, ["Gunakan peta sebagai daftar awal audit pascakebakaran, restorasi, dan pengecekan akses pemadaman.", "Pisahkan anggaran pemulihan jejak 2024 dari anggaran kesiapsiagaan musim mendatang.", "Overlay dengan izin, aset, dan tutupan lahan tervalidasi sebelum menetapkan pemilik tindakan atau nilai kerugian."])
    doc.add_heading("Batasan interpretasi", 1)
    doc.add_paragraph("Produk annual memperlihatkan perubahan permukaan yang diasosiasikan dengan pembakaran, bukan waktu penyalaan, besar api, ataupun penyebab kebakaran. Produk monitoring 2026 saat ini hanya Januari–Juni dan pangsa bulanan dapat memiliki overlap antarbulan. Jangan menjumlahkan nilai bulanan atau memperlakukan jumlah piksel sebagai hektare tanpa metodologi area yang terdokumentasi.")
    finish(doc, output, "03_MapBiomas_Burned_Area_EDA.docx")


def peat_report(output: Path, figures: Path) -> None:
    provinces, province_col = kalimantan(); peat = gpd.read_file(ROOT / "data/raw/peatland/kalimantan_peatlands.geojson")
    area = provinces.to_crs("EPSG:6933").set_index(province_col).area / 1e6
    intersection = gpd.overlay(provinces[[province_col, "geometry"]], peat[["geometry"]], how="intersection").to_crs("EPSG:6933")
    intersection["area_km2"] = intersection.geometry.area / 1e6
    peat_area = intersection.groupby(province_col)["area_km2"].sum()
    summary = pd.DataFrame({"province_area_km2": area, "peat_km2": peat_area}).fillna(0); summary["peat_share_pct"] = summary.peat_km2 / summary.province_area_km2 * 100
    summary = summary.sort_values("peat_km2", ascending=False); top = summary.iloc[0]
    highest_dependency = summary.sort_values("peat_share_pct", ascending=False).iloc[0]
    plt.figure(figsize=(9.4, 4.7)); plt.barh(summary.index, summary.peat_km2, color=GREEN); plt.gca().invert_yaxis(); plt.xlabel("Luas poligon gambut (km², equal-area)"); plt.title("Cakupan gambut per provinsi Kalimantan"); plt.grid(axis="x", alpha=.25)
    save_plot(figures / "peat_bars.png")
    fig, ax = plt.subplots(figsize=(7, 5.3)); provinces.to_crs("EPSG:4326").plot(ax=ax, color="#D9EAF7", edgecolor="white"); peat.to_crs("EPSG:4326").plot(ax=ax, color="#8C6D31", alpha=.65, linewidth=.1); ax.set_title("Sebaran poligon gambut Kalimantan"); ax.set_axis_off()
    save_plot(figures / "peat_map.png")
    fragmentation = intersection.groupby(province_col).agg(polygon_count=("area_km2", "size"), median_polygon_km2=("area_km2", "median"))
    plt.figure(figsize=(7.5, 5.1))
    for name, row in fragmentation.iterrows(): plt.scatter(row.polygon_count, row.median_polygon_km2, s=80, color=GREEN); plt.annotate(name, (row.polygon_count, row.median_polygon_km2), xytext=(5, 4), textcoords="offset points", fontsize=8)
    plt.xlabel("Jumlah fragmen poligon gambut"); plt.ylabel("Median luas fragmen (km²)"); plt.title("Fragmentasi gambut dan kompleksitas implementasi")
    plt.grid(alpha=.25)
    save_plot(figures / "peat_fragmentation.png")
    peat_dependency = summary.sort_values("peat_share_pct")
    plt.figure(figsize=(9.1, 4.6)); plt.barh(peat_dependency.index, peat_dependency.peat_share_pct, color="#6D597A")
    plt.xlabel("Pangsa gambut terhadap luas provinsi (%)"); plt.title("Ketergantungan gambut relatif per provinsi")
    plt.grid(axis="x", alpha=.25)
    save_plot(figures / "peat_province_share.png")
    priority = summary.join(fragmentation, how="left")
    plt.figure(figsize=(7.7, 5.2))
    for name, row in priority.iterrows():
        plt.scatter(row.peat_share_pct, row.peat_km2, s=35 + row.polygon_count * 2.4, color="#8C6D31", alpha=.75)
        plt.annotate(name, (row.peat_share_pct, row.peat_km2), xytext=(5, 4), textcoords="offset points", fontsize=8)
    plt.xlabel("Pangsa gambut dalam provinsi (%)"); plt.ylabel("Luas gambut (km²)")
    plt.title("Kuadran skala dan ketergantungan gambut per provinsi")
    plt.grid(alpha=.25)
    save_plot(figures / "peat_province_priority.png")
    doc = document("WRI/GFW: gambut dan prioritas pengelolaan air", "EDA bisnis Kalimantan | tanpa NASA FIRMS maupun data populasi")
    doc.add_heading("Tentang dataset", 1)
    doc.add_paragraph(f"Dataset WRI/Global Forest Watch di proyek berisi {len(peat):,} poligon area gambut yang telah dipotong tepat terhadap gabungan lima provinsi Kalimantan. Ekstraksi awal dari endpoint WRI/GFW menghasilkan 350 poligon dalam bounding box; satu geometri di luar pulau dieliminasi oleh pemotongan geometri sehingga keluaran ini tidak lagi hanya subset kotak koordinat. Poligon diukur pada proyeksi equal-area EPSG:6933 agar perbandingan luas tidak bias oleh koordinat geografis. Hasil adalah luas geometri sumber, bukan delineasi legal terbaru atau pengukuran kedalaman gambut.")
    feature_table(doc, [
        ("objectid_1", "Nomor ID", "Nomor urut poligon dari satu versi ekspor data", "Hanya untuk mengecek duplikasi; bukan indikator gambut atau kebakaran"),
        ("objectid", "Nomor ID", "Nomor ID lain untuk poligon yang sama", "Redundan dengan objectid_1 pada file ini; hanya untuk pelacakan"),
        ("globalid", "Kode unik (UUID)", "Identitas unik poligon yang paling aman untuk dirujuk lintas ekspor", "Dipakai bila perlu menelusuri poligon tertentu; bukan fitur model"),
        ("layer_revi", "Label teks", "Pada hampir semua baris berisi 'gambut'—penanda layer peta", "Bukan tahun, bukan status kondisi gambut, dan bukan tingkat risiko"),
        ("shape_leng", "Angka metadata", "Perhitungan panjang batas poligon dari sistem sumber", "Tidak dipakai sebagai panjang kanal atau jarak operasi"),
        ("shape_Length", "Angka metadata", "Perhitungan panjang batas poligon versi ArcGIS", "Mirip shape_leng tetapi tidak identik; tidak dipakai dalam analisis"),
        ("shape_Area", "Angka metadata", "Luas poligon sebagaimana tersimpan di sistem sumber", "Tidak dipakai langsung; luas dihitung ulang dari geometri dalam proyeksi yang jelas"),
        ("geometry", "Bentuk/lokasi poligon", "Garis batas area gambut di peta", "Kolom paling penting: dioverlay dengan provinsi dan grid untuk menghitung peat_pct"),
        ("tanggal/tahun", "Tidak ada", "Layer ini tidak menyimpan waktu pengamatan", "Baseline peta tetap untuk 2023–2026; tidak boleh dianggap tren tahunan"),
    ])
    metrics(doc, [("Poligon setelah clip Kalimantan", f"{len(peat):,} dari 350 hasil query bounding box"), ("Cakupan waktu", "Layer statis; dapat dipakai sebagai baseline 2023–2026, bukan data tahunan"), ("Luas gambut ter-overlay", f"{summary.peat_km2.sum():,.0f} km²"), ("Provinsi dengan luas gambut terbesar", f"{top.name} ({top.peat_km2:,.0f} km²)"), ("Ketergantungan gambut tertinggi", f"{highest_dependency.name} ({highest_dependency.peat_share_pct:.1f}% luas provinsi)")])
    doc.add_heading("Exploratory Data Analysis", 1)
    chart(doc, figures / "peat_bars.png", "Gambar 1. Luas adalah irisan poligon gambut dan batas provinsi pada proyeksi equal-area.")
    doc.add_paragraph(f"{top.name} memegang cakupan geometri gambut terbesar ({top.peat_km2:,.0f} km²), sedangkan {highest_dependency.name} memiliki ketergantungan relatif tertinggi ({highest_dependency.peat_share_pct:.1f}% dari luas provinsi). Dua peringkat ini dapat menghasilkan prioritas investasi yang berbeda: skala absolut untuk kapasitas regional dan porsi relatif untuk intensitas tata kelola lokal.")
    chart(doc, figures / "peat_map.png", "Gambar 2. Geometri gambut terhadap batas provinsi untuk perencanaan koridor pengelolaan air.")
    chart(doc, figures / "peat_fragmentation.png", "Gambar 3. Banyak fragmen dan ukuran median membantu memilih model implementasi pengelolaan air.")
    doc.add_paragraph("EDA fragmentasi membedakan kebutuhan implementasi. Banyak fragmen kecil biasanya memerlukan lebih banyak titik inspeksi, relasi lokal, dan pemantauan tersebar. Fragmen yang lebih besar mendukung pendekatan koridor atau proyek hidrologi skala lanskap. Grafik ini tidak mengukur konektivitas kanal, sehingga harus dilengkapi data hidrologi sebelum desain final.")
    chart(doc, figures / "peat_province_share.png", "Gambar 4. Ketergantungan relatif: proporsi area provinsi yang beririsan dengan poligon gambut, bukan kondisi basah atau kedalaman gambut.")
    doc.add_paragraph(f"Grafik proporsi memperlihatkan alasan {highest_dependency.name} perlu mendapat prosedur operasi gambut yang kuat meskipun skala absolutnya bukan yang terbesar. Ukuran investasi regional dan intensitas tata kelola lokal harus dibedakan agar provinsi kecil tetapi bergambut tinggi tidak terlewat.")
    chart(doc, figures / "peat_province_priority.png", "Gambar 5. Kuadran per provinsi: sumbu-X adalah ketergantungan relatif, sumbu-Y adalah skala gambut; ukuran titik mencerminkan jumlah fragmen hasil overlay.")
    doc.add_paragraph("Kuadran memudahkan pembagian portofolio: posisi tinggi berarti kebutuhan kapasitas program besar; posisi ke kanan berarti kebutuhan SOP dan pendampingan gambut lebih intensif; titik besar menandakan kompleksitas titik inspeksi yang lebih tinggi. Ini adalah dasar pemilihan tipe intervensi, bukan skor legalitas atau risiko api saat ini.")
    doc.add_heading("Implikasi terhadap karhutla Kalimantan", 1)
    doc.add_paragraph("Pada gambut, masalah karhutla adalah masalah pengelolaan air sebelum menjadi masalah pemadaman. Ketika muka air turun, api dapat menjalar di bawah permukaan dan membutuhkan operasi yang lebih lama daripada api permukaan. Karena itu, ukuran keberhasilan awal yang lebih berguna adalah kesiapan hidrologi—kondisi kanal, muka air, dan pemeliharaan sekat—bukan hanya jumlah sumber daya yang dikerahkan setelah kejadian.")
    insight(doc, "Kelola gambut sebagai portofolio hidrologi lintas batas. Gunakan luasan besar untuk menentukan kapasitas regional, porsi gambut tinggi untuk menentukan intensitas pendampingan lokal, dan kesatuan hidrologis untuk menentukan lokasi sekat kanal, pemantauan muka air, serta kemitraan desa.")
    doc.add_heading("Rekomendasi keputusan", 1)
    bullets(doc, ["Buat daftar proyek pengelolaan air dari klaster poligon gambut yang berdekatan dan lintas batas.", "Gunakan luas sebagai ukuran awal skala investasi lalu validasi kedalaman gambut, kanal, akses, dan pengelolaan.", "Ukur manfaat investasi dengan indikator muka air, kondisi sekat, dan kepatuhan pemeliharaan."])
    doc.add_heading("Batasan interpretasi", 1)
    doc.add_paragraph("Peta ini tidak membuktikan kondisi basah/kering saat ini dan tidak cukup untuk menetapkan kewajiban legal. Layer tidak memiliki atribut tahun/tanggal, sehingga tidak boleh dipakai untuk menyimpulkan perubahan gambut antar-2023, 2024, 2025, dan 2026. Dalam model, ia hanya sah sebagai kovariat struktural yang nilainya tetap selama periode tersebut. Validasi hidrologi dan lapangan diperlukan sebelum desain intervensi.")
    finish(doc, output, "04_WRI_GFW_Peatlands_EDA.docx")


def boundaries_report(output: Path, figures: Path) -> None:
    provinces, province_col = kalimantan(); adm2 = gpd.read_file(ROOT / "data/raw/boundaries/idn_adm2.geojson")
    province_key = "province_name"
    province_shapes = provinces[[province_col, "geometry"]].rename(columns={province_col: province_key})
    districts = gpd.sjoin(adm2.to_crs(provinces.crs), province_shapes, how="inner", predicate="intersects")
    if "shapeID" in districts.columns: districts = districts.drop_duplicates("shapeID")
    area = provinces.to_crs("EPSG:6933").set_index(province_col).area / 1e6; area.index.name = province_key
    count = districts.groupby(province_key).size()
    summary = pd.DataFrame({"province_area_km2": area, "district_count": count}).fillna(0); summary["area_per_district_km2"] = summary.province_area_km2 / summary.district_count.replace(0, np.nan)
    summary = summary.sort_values("province_area_km2", ascending=False); largest = summary.iloc[0]; most = summary.sort_values("district_count", ascending=False).iloc[0]
    widest_territory = summary.sort_values("area_per_district_km2", ascending=False).iloc[0]
    fig, ax1 = plt.subplots(figsize=(9.6, 4.7)); x = np.arange(len(summary)); width = .38
    ax1.bar(x-width/2, summary.province_area_km2, width, color=BLUE, label="Luas provinsi (km²)"); ax2 = ax1.twinx(); ax2.bar(x+width/2, summary.district_count, width, color=ORANGE, label="Jumlah ADM2")
    ax1.set_xticks(x, summary.index, rotation=18, ha="right"); ax1.set_ylabel("Luas (km²)"); ax2.set_ylabel("Jumlah kabupaten/kota"); ax1.set_title("Jangkauan fisik dan beban koordinasi administratif")
    a, b = ax1.get_legend_handles_labels(); c, d = ax2.get_legend_handles_labels(); ax1.legend(a+c, b+d, loc="upper right")
    save_plot(figures / "admin_bars.png")
    fig, ax = plt.subplots(figsize=(7, 5.3)); provinces.to_crs("EPSG:4326").plot(ax=ax, color="#D9EAF7", edgecolor="white"); districts.to_crs("EPSG:4326").boundary.plot(ax=ax, color="#4A4A4A", linewidth=.22); ax.set_title("Batas provinsi dan kabupaten/kota Kalimantan"); ax.set_axis_off()
    save_plot(figures / "admin_map.png")
    coverage_load = summary.sort_values("area_per_district_km2")
    plt.figure(figsize=(8.6, 4.8)); plt.barh(coverage_load.index, coverage_load.area_per_district_km2, color=BLUE)
    plt.xlabel("Rata-rata luas per kabupaten/kota (km²)"); plt.title("Beban jangkauan rata-rata per ADM2")
    plt.grid(axis="x", alpha=.25)
    save_plot(figures / "admin_territory_per_district.png")
    summary["districts_per_10000_km2"] = summary.district_count / summary.province_area_km2 * 10_000
    coordination_density = summary.sort_values("districts_per_10000_km2")
    plt.figure(figsize=(8.8, 4.7)); plt.barh(coordination_density.index, coordination_density.districts_per_10000_km2, color=ORANGE)
    plt.xlabel("Kabupaten/kota per 10.000 km²"); plt.title("Kepadatan titik koordinasi per provinsi")
    plt.grid(axis="x", alpha=.25)
    save_plot(figures / "admin_coordination_density.png")
    plt.figure(figsize=(7.6, 5.1))
    for name, row in summary.iterrows():
        plt.scatter(row.area_per_district_km2, row.district_count, s=45 + row.province_area_km2 / 1_000, color=BLUE, alpha=.75)
        plt.annotate(name, (row.area_per_district_km2, row.district_count), xytext=(5, 4), textcoords="offset points", fontsize=8)
    plt.xlabel("Rata-rata luas per ADM2 (km²)"); plt.ylabel("Jumlah kabupaten/kota")
    plt.title("Kuadran jangkauan lapangan dan kompleksitas koordinasi")
    plt.grid(alpha=.25)
    save_plot(figures / "admin_province_operating_quadrant.png")
    densest = summary.sort_values("districts_per_10000_km2", ascending=False).iloc[0]
    doc = document("geoBoundaries: cakupan operasi dan koordinasi", "EDA bisnis Kalimantan | tanpa NASA FIRMS maupun data populasi")
    doc.add_heading("Tentang dataset", 1)
    doc.add_paragraph("geoBoundaries menyediakan geometri administrasi terbuka: 34 fitur ADM1 nasional dan 519 fitur ADM2 nasional pada snapshot sumber proyek. Analisis memilih provinsi bernama Kalimantan lalu mengaitkan ADM2 melalui irisan spasial, bukan asumsi nilai kolom. Luas dihitung dalam proyeksi equal-area EPSG:6933.")
    feature_table(doc, [
        ("shapeName", "Nama teks", "Nama wilayah administrasi, misalnya nama provinsi atau kabupaten/kota", "Memilih lima provinsi Kalimantan dan memberi label grafik"),
        ("shapeISO", "Kode wilayah", "Kode identitas administratif yang disediakan sumber", "Membantu menelusuri wilayah; bukan skor atau indikator risiko"),
        ("shapeID", "Kode unik", "Nomor/ID unik untuk satu batas wilayah", "Mencegah satu kabupaten/kota terhitung dua kali saat overlay"),
        ("shapeGroup", "Nama negara", "Kelompok negara dari geoBoundaries", "Memastikan data memang konteks Indonesia; bukan nama provinsi"),
        ("shapeType", "Label level", "Menandakan level batas: ADM1 = provinsi, ADM2 = kabupaten/kota", "Membedakan pihak koordinasi tingkat provinsi dan tingkat kabupaten/kota"),
        ("geometry", "Garis/bentuk peta", "Bentuk batas wilayah di peta", "Dasar menghitung luas, overlay, pembagian area kerja, dan peta"),
        ("Level dataset", "Cakupan data", "Sumber menyediakan 34 ADM1 dan 519 ADM2 Indonesia", "Dipilih hanya yang beririsan dengan Kalimantan untuk membaca beban logistik dan koordinasi"),
    ])
    metrics(doc, [("Provinsi Kalimantan", str(len(provinces))), ("Kabupaten/kota ter-overlay", f"{len(districts):,}"), ("Provinsi terluas", f"{largest.name} ({largest.province_area_km2:,.0f} km²)"), ("Koordinasi ADM2 tertinggi", f"{most.name} ({int(most.district_count)} kabupaten/kota)"), ("Jangkauan rata-rata per ADM2 terluas", f"{widest_territory.name} ({widest_territory.area_per_district_km2:,.0f} km²/ADM2)")])
    doc.add_heading("Exploratory Data Analysis", 1)
    chart(doc, figures / "admin_bars.png", "Gambar 1. Luas mengukur jangkauan fisik; ADM2 mengukur banyaknya mitra koordinasi potensial.")
    doc.add_paragraph(f"{largest.name} memiliki cakupan fisik terbesar, sedangkan {most.name} memiliki jumlah kabupaten/kota terbanyak. Keduanya menjawab kebutuhan yang berbeda: perjalanan/logistik versus rapat, pelaporan, dan penyelarasan pemangku kepentingan.")
    chart(doc, figures / "admin_map.png", "Gambar 2. Struktur administrasi untuk membagi wilayah kerja dan rute koordinasi.")
    chart(doc, figures / "admin_territory_per_district.png", "Gambar 3. Rata-rata luas per ADM2 sebagai proksi beban jangkauan lapangan, bukan tingkat risiko.")
    doc.add_paragraph(f"{widest_territory.name} memiliki rata-rata jangkauan per ADM2 terbesar ({widest_territory.area_per_district_km2:,.0f} km²). Pada konfigurasi seperti ini, satu pola posko atau satu jadwal patroli seragam berisiko menciptakan waktu tempuh yang terlalu panjang. Ini adalah dasar untuk merancang hub-and-spoke, titik logistik, atau dukungan bergerak sebelum memasukkan data risiko kebakaran.")
    chart(doc, figures / "admin_coordination_density.png", "Gambar 4. Kepadatan titik koordinasi: jumlah kabupaten/kota relatif terhadap luas provinsi, bukan ukuran kapasitas atau penduduk.")
    doc.add_paragraph(f"{densest.name} memiliki kepadatan unit ADM2 tertinggi ({densest.districts_per_10000_km2:.2f} kabupaten/kota per 10.000 km²). Di wilayah seperti ini, liaison, alur pelaporan, dan daftar kontak perlu diskalakan lebih cepat daripada logistik fisik semata—karena keputusan dapat tertunda oleh banyaknya pihak yang harus disejajarkan.")
    chart(doc, figures / "admin_province_operating_quadrant.png", "Gambar 5. Kuadran per provinsi: kanan berarti jangkauan lapangan lebih lebar per ADM2; atas berarti lebih banyak mitra administrasi yang perlu dikoordinasikan.")
    doc.add_paragraph("Kuadran menghindari kesalahan menyamakan provinsi besar dengan provinsi kompleks secara administratif. Titik kanan-atas perlu dua investasi sekaligus: hub logistik yang mampu menjangkau area luas dan kapasitas liaison/manajemen informasi yang mampu menyatukan banyak pemda. Titik lain dapat mengutamakan salah satunya.")
    doc.add_heading("Implikasi terhadap karhutla Kalimantan", 1)
    doc.add_paragraph("Operasi karhutla memiliki dua hambatan yang berbeda: waktu mencapai lokasi dan waktu menyatukan keputusan. Rasio luas per ADM2 menunjukkan beban jangkauan rata-rata; jumlah ADM2 menunjukkan banyaknya titik koordinasi untuk status siaga, izin akses, komunikasi masyarakat, dan pelaporan. Command structure yang hanya mengoptimalkan salah satunya akan menimbulkan keterlambatan pada yang lain.")
    insight(doc, "Rancang struktur respons dengan dua kapasitas: tim lapangan/pos logistik untuk wilayah dengan jangkauan per ADM2 besar, serta liaison dan manajemen informasi untuk wilayah dengan banyak ADM2. Saat indikator cuaca atau gambut menaikkan status siaga, struktur ini memungkinkan eskalasi pencegahan sebelum peristiwa menjadi lintas kabupaten.")
    doc.add_heading("Rekomendasi keputusan", 1)
    bullets(doc, ["Rancang wilayah kerja lapangan dari luas dan akses, lalu peran liaison dari jumlah pemda yang dilibatkan.", "Gunakan batas ADM2 untuk mencegah area kosong dan tumpang tindih pelaporan.", "Verifikasi batas administrasi terbaru sebelum keputusan formal karena geoBoundaries adalah referensi terbuka, bukan penetapan legal."])
    doc.add_heading("Batasan interpretasi", 1)
    doc.add_paragraph("Jumlah ADM2 bukan ukuran penduduk, anggaran, tingkat risiko, maupun kapasitas pemda. Data ini hanya menggambarkan struktur ruang dan kompleksitas koordinasi potensial.")
    finish(doc, output, "05_geoBoundaries_Administrative_Coverage_EDA.docx")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--only",
        choices=("era5", "landcover", "burned", "peat", "boundaries"),
        help="Generate one selected report; omit to generate all five.",
    )
    args = parser.parse_args()
    output = args.output
    tasks = [
        ("era5", era5_report),
        ("landcover", landcover_report),
        ("burned", burned_report),
        ("peat", peat_report),
        ("boundaries", boundaries_report),
    ]
    with tempfile.TemporaryDirectory(prefix="fireline_eda_") as directory:
        figures = Path(directory)
        for name, task in tasks:
            if args.only and args.only != name:
                continue
            print(f"starting {name}", flush=True)
            task(output, figures)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
